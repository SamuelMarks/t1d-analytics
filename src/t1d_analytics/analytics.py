"""Analytics module for querying downloaded T1D datasets."""

import csv
import hashlib
import json
import os
import re
import tarfile
import zipfile
from pathlib import Path
from typing import TYPE_CHECKING, Any, Dict, List, Optional, Union

import duckdb

if TYPE_CHECKING:
    import pandas as pd  # type: ignore[import-untyped]

from t1d_analytics.diagnostics import DatabaseStatusCode, check_database_health
from t1d_analytics.i18n import get_translator


def compute_file_hash(path: Path) -> str:
    """
    Compute SHA-256 hash of a file.

    Args:
    ----
        path: Path to the file.

    Returns:
    -------
        str: Hexadecimal SHA-256 digest of the file contents.

    """
    h = hashlib.sha256()
    with open(path, "rb") as f:
        while True:
            chunk = f.read(65536)
            if not chunk:
                break
            h.update(chunk)
    return h.hexdigest()


def extract_zips(data_dir_str: str, max_depth: int = 3) -> None:
    """
    Extract zip and tar archives found in the data directory safely without path traversal.

    Recursively decompresses nested archives up to max_depth.

    Args:
    ----
        data_dir_str: Directory path containing archive files to extract.
        max_depth: Maximum recursion depth for decompressing nested archives.

    """
    _ = get_translator()
    data_dir = Path(data_dir_str)
    if not data_dir.exists():
        print(_("Data directory {} does not exist.", data_dir))
        return

    print(_("Checking for zip files to extract..."))
    processed_archives: set[Path] = set()
    found_any = False
    archive_patterns = ("*.zip", "*.tar.gz", "*.tgz", "*.tar.bz2", "*.tar")

    for depth in range(max_depth):
        candidate_archives: list[Path] = []
        for pat in archive_patterns:
            candidate_archives.extend(data_dir.rglob(pat))
        candidate_archives = list(dict.fromkeys(candidate_archives))

        unprocessed = [
            p for p in candidate_archives if p.resolve() not in processed_archives
        ]
        if not unprocessed:
            break

        for path in unprocessed:
            found_any = True
            resolved_archive = path.resolve()
            processed_archives.add(resolved_archive)

            # Determine extraction directory
            p_name = path.name.lower()
            if p_name.endswith(".tar.gz") or p_name.endswith(".tar.bz2"):
                extract_stem = path.name[:-7]
                extract_dir = path.parent / extract_stem
            elif p_name.endswith(".tgz"):
                extract_stem = path.name[:-4]
                extract_dir = path.parent / extract_stem
            else:
                extract_dir = path.with_suffix("")

            if depth == 0 and extract_dir.exists():
                continue

            print(_("Extracting {}...", path.name))
            extract_dir.mkdir(parents=True, exist_ok=True)
            resolved_extract_dir = extract_dir.resolve()

            is_tar = any(
                p_name.endswith(ext) for ext in (".tar.gz", ".tgz", ".tar.bz2", ".tar")
            )
            if is_tar:
                try:
                    with tarfile.open(path, "r:*") as tar_ref:
                        for tar_member in tar_ref.getmembers():
                            if tar_member.issym() or tar_member.islnk():
                                print(
                                    _(
                                        "Skipping unsafe entry {} in {} (symlink or hardlink detected).",
                                        tar_member.name,
                                        path.name,
                                    )
                                )
                                continue
                            tar_target = (
                                resolved_extract_dir / tar_member.name
                            ).resolve()
                            is_safe_tar = (
                                tar_target == resolved_extract_dir
                                or tar_target.is_relative_to(resolved_extract_dir)
                            )
                            if not is_safe_tar:
                                print(
                                    _(
                                        "Skipping unsafe entry {} in {} (path traversal detected).",
                                        tar_member.name,
                                        path.name,
                                    )
                                )
                                continue
                            tar_ref.extract(tar_member, extract_dir)
                except (tarfile.TarError, EOFError, OSError):
                    print(_("Failed to extract {} (Bad Archive File).", path.name))
            else:
                try:
                    with zipfile.ZipFile(path, "r") as zip_ref:
                        for zip_info in zip_ref.infolist():
                            # Reject zip symlinks (mode 0o120000)
                            if (zip_info.external_attr >> 16) & 0o170000 == 0o120000:
                                print(
                                    _(
                                        "Skipping unsafe entry {} in {} (symlink detected).",
                                        zip_info.filename,
                                        path.name,
                                    )
                                )
                                continue
                            zip_target = (
                                resolved_extract_dir / zip_info.filename
                            ).resolve()
                            is_safe_zip = (
                                zip_target == resolved_extract_dir
                                or zip_target.is_relative_to(resolved_extract_dir)
                            )
                            if not is_safe_zip:
                                print(
                                    _(
                                        "Skipping unsafe entry {} in {} (path traversal detected).",
                                        zip_info.filename,
                                        path.name,
                                    )
                                )
                                continue
                            zip_ref.extract(zip_info, extract_dir)
                except zipfile.BadZipFile:
                    print(_("Failed to extract {} (Bad Zip File).", path.name))

    if not found_any:
        print(_("No zip files found to extract."))
    else:
        print(_("Extraction complete."))


def _read_clinical_dataframe(data_file: Path) -> "pd.DataFrame":
    """
    Read SAS (.sas7bdat, .xpt) or SPSS (.sav) data file into a pandas DataFrame.

    Translates SAS missing values and dates to standard nullable types.

    Args:
    ----
        data_file: Path to the clinical data file.

    Returns:
    -------
        pd.DataFrame: Parsed DataFrame with converted column data.

    """
    import pandas as pd

    suffix = data_file.suffix.lower()
    df: Optional[pd.DataFrame] = None
    if suffix in (".sas7bdat", ".xpt"):
        try:
            import pyreadstat  # type: ignore[import-untyped]

            if suffix == ".sas7bdat":
                df, _ = pyreadstat.read_sas7bdat(str(data_file))
            else:
                df, _ = pyreadstat.read_xport(str(data_file))
        except (ImportError, Exception):
            fmt = "sas7bdat" if suffix == ".sas7bdat" else "xport"
            df = pd.read_sas(str(data_file), format=fmt)
    elif suffix == ".sav":
        try:
            import pyreadstat

            df, _ = pyreadstat.read_sav(str(data_file))
        except (ImportError, Exception):
            df = pd.read_spss(str(data_file))
    else:
        raise ValueError(f"Unsupported clinical file format: {suffix}")

    if df is None:
        raise ValueError(f"Could not read clinical file {data_file}")

    # Ensure column names are standard strings
    df.columns = [
        col.decode("utf-8", errors="ignore") if isinstance(col, bytes) else str(col)
        for col in df.columns
    ]
    # Decode byte string elements if any
    for col in df.columns:
        if df[col].dtype == object:
            df[col] = df[col].apply(
                lambda x: (
                    x.decode("utf-8", errors="ignore") if isinstance(x, bytes) else x
                )
            )
    return df


def get_ingestion_manifest(db_path: str) -> list[dict[str, object]]:
    """
    Retrieve the ingestion manifest audit records from DuckDB.

    Args:
    ----
        db_path: Path to the DuckDB database file.

    Returns:
    -------
        list[dict[str, object]]: List of manifest records containing source_file,
            table_name, file_hash, row_count, loaded_at, and file_size.

    """
    if not Path(db_path).exists():
        return []

    conn = duckdb.connect(db_path, read_only=True)
    has_manifest = conn.execute(
        "SELECT 1 FROM information_schema.tables WHERE table_name = '_t1d_ingestion_manifest'"
    ).fetchone()
    if not has_manifest:
        conn.close()
        return []

    rows = conn.execute(
        "SELECT source_file, table_name, file_hash, row_count, loaded_at, file_size FROM _t1d_ingestion_manifest ORDER BY loaded_at"
    ).fetchall()
    conn.close()

    return [
        {
            "source_file": r[0],
            "table_name": r[1],
            "file_hash": r[2],
            "row_count": r[3],
            "loaded_at": str(r[4]),
            "file_size": r[5],
        }
        for r in rows
    ]


def get_clinical_variable_mappings() -> Dict[str, List[str]]:
    """
    Load the clinical variable ontology mapping standard concepts to trial-specific column names.

    Returns
    -------
        Dict[str, List[str]]: Mapping from standard concept names to list of original column names.

    """
    matches_file = Path(__file__).parent / "likely_matches.json"
    if matches_file.exists():
        try:
            with open(matches_file, "r") as f:
                data = json.load(f)
                if isinstance(data, dict):
                    return {
                        str(k): [str(x) for x in v]
                        for k, v in data.items()
                        if isinstance(v, list)
                    }
        except Exception:
            pass
    return {}


def load_data_to_duckdb(
    data_dir_str: str,
    db_path: str,
    prefix_subdirs: bool = False,
    include_sas: bool = False,
    include_spss: bool = False,
) -> None:
    """
    Load tabular files from the data directory into DuckDB tables.

    Supports CSV, TXT, Parquet, GZ, and optionally SAS and SPSS datasets.
    Tracks ingested file hashes in `_t1d_ingestion_manifest` to prevent redundant ingestion.

    Args:
    ----
        data_dir_str: Directory containing extracted CSV/TXT datasets.
        db_path: Path to the target DuckDB database file.
        prefix_subdirs: Whether to prefix table names with parent subdirectories to disambiguate.
        include_sas: Whether to include SAS clinical files (.sas7bdat, .xpt).
        include_spss: Whether to include SPSS clinical files (.sav).

    """
    _ = get_translator()
    data_dir = Path(data_dir_str)
    if not data_dir.exists():
        print(_("Data directory {} does not exist.", data_dir))
        return

    print(_("Connecting to DuckDB at {}...", db_path))
    conn = duckdb.connect(db_path, read_only=False)
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS _t1d_ingestion_manifest (
            source_file VARCHAR,
            table_name VARCHAR,
            file_hash VARCHAR,
            row_count BIGINT,
            loaded_at TIMESTAMP,
            file_size BIGINT
        )
        """
    )

    print(_("Scanning for tabular files (CSV, TXT, Parquet, GZ)..."))
    patterns = [
        "*.csv",
        "*.CSV",
        "*.txt",
        "*.TXT",
        "*.parquet",
        "*.PARQUET",
        "*.csv.gz",
        "*.CSV.GZ",
    ]
    if include_sas:
        patterns.extend(["*.sas7bdat", "*.SAS7BDAT", "*.xpt", "*.XPT"])
    if include_spss:
        patterns.extend(["*.sav", "*.SAV"])

    raw_files: list[Path] = []
    for pat in patterns:
        raw_files.extend(data_dir.rglob(pat))
    data_files = list(dict.fromkeys(raw_files))

    # Filter out MacOS hidden files or unhelpful docs
    data_files = [
        f
        for f in data_files
        if not f.name.startswith("._")
        and "Glossary" not in f.name
        and "Protocol" not in f.name
    ]

    if not data_files:
        print(_("No tabular files found."))
        conn.close()
        return

    # Load reverse map for column renaming
    reverse_map: Dict[str, str] = {}
    matches = get_clinical_variable_mappings()
    for std_name, orig_names in matches.items():
        for orig in orig_names:
            reverse_map[orig.lower()] = std_name

    stem_counts: dict[str, int] = {}
    for df in data_files:
        s_name = Path(df.stem).stem if df.name.endswith(".csv.gz") else df.stem
        raw_stem = re.sub(r"[^a-zA-Z0-9_]", "_", s_name).lower().strip("_") or "dataset"
        stem_counts[raw_stem] = stem_counts.get(raw_stem, 0) + 1

    existing_tables = [
        t[0]
        for t in conn.execute("SHOW TABLES").fetchall()
        if not t[0].startswith("_t1d_")
    ]
    loaded_tables = list(existing_tables)
    for data_file in data_files:
        s_name = (
            Path(data_file.stem).stem
            if data_file.name.endswith(".csv.gz")
            else data_file.stem
        )
        raw_stem = re.sub(r"[^a-zA-Z0-9_]", "_", s_name).lower().strip("_") or "dataset"
        safe_stem = raw_stem
        if prefix_subdirs or stem_counts.get(raw_stem, 0) > 1:
            rel_parent = data_file.parent.relative_to(data_dir)
            if rel_parent != Path("."):
                prefix = (
                    re.sub(r"[^a-zA-Z0-9_]", "_", str(rel_parent)).strip("_").lower()
                )
                safe_stem = f"{prefix}_{raw_stem}" if prefix else raw_stem

        current_source = str(data_file.resolve())
        table_name = safe_stem.strip("_") or "dataset"
        # Ensure it starts with a letter
        if not table_name[0].isalpha():
            table_name = "t_" + table_name

        counter = 1
        base_tbl = table_name
        is_already_loaded = False

        try:
            current_hash = compute_file_hash(data_file)
            manifest_row = conn.execute(
                "SELECT table_name FROM _t1d_ingestion_manifest WHERE source_file = ? AND file_hash = ?",
                [current_source, current_hash],
            ).fetchone()

            while table_name in existing_tables or table_name in loaded_tables:
                row = conn.execute(
                    "SELECT comment FROM duckdb_tables() WHERE table_name = ?",
                    [table_name],
                ).fetchone()
                existing_comment = row[0] if row else None

                if existing_comment == current_source or (
                    manifest_row and manifest_row[0] == table_name
                ):
                    is_already_loaded = True
                    break

                counter += 1
                table_name = f"{base_tbl}_{counter}"

            if is_already_loaded:
                print(_("Table {} already exists, skipping...", table_name))
                continue

            escaped_path = str(data_file).replace("'", "''")
            is_clinical = data_file.suffix.lower() in (
                ".sas7bdat",
                ".xpt",
                ".sav",
            )
            if data_file.suffix.lower() == ".parquet":
                print(
                    _(
                        "Loading {} (parquet) into table {}...",
                        data_file.name,
                        table_name,
                    )
                )
                conn.execute(
                    f"CREATE OR REPLACE TEMP VIEW temp_view AS SELECT * FROM read_parquet('{escaped_path}')"
                )
            elif data_file.name.endswith(".csv.gz"):
                print(
                    _(
                        "Loading {} (csv.gz) into table {}...",
                        data_file.name,
                        table_name,
                    )
                )
                conn.execute(
                    f"CREATE OR REPLACE TEMP VIEW temp_view AS SELECT * FROM read_csv('{escaped_path}', auto_detect=true)"
                )
            elif is_clinical:
                print(
                    _(
                        "Loading {} (clinical) into table {}...",
                        data_file.name,
                        table_name,
                    )
                )
                clinical_df = _read_clinical_dataframe(data_file)
                conn.register("_clinical_df", clinical_df)
                conn.execute(
                    "CREATE OR REPLACE TEMP VIEW temp_view AS SELECT * FROM _clinical_df"
                )
            else:
                # 1. Detect Encoding (utf-8-sig vs utf-16 vs utf-8)
                with open(data_file, "rb") as f_rb:
                    raw_bytes = f_rb.read(4)
                if raw_bytes[:2] in (b"\xff\xfe", b"\xfe\xff"):
                    py_encoding = "utf-16"
                    duckdb_encoding = "utf-16"
                elif raw_bytes[:3] == b"\xef\xbb\xbf":
                    py_encoding = "utf-8-sig"
                    duckdb_encoding = "utf-8"
                else:
                    py_encoding = "utf-8"
                    duckdb_encoding = "utf-8"

                # 2. Detect Separator using csv.Sniffer with fallback
                sep = ","
                try:
                    with open(
                        data_file, "r", encoding=py_encoding, errors="ignore"
                    ) as f:
                        sample = "".join([f.readline() for _ in range(5)])
                    if sample.strip():
                        dialect = csv.Sniffer().sniff(sample, delimiters=",\t|;")
                        sep = dialect.delimiter
                except Exception:
                    with open(
                        data_file, "r", encoding=py_encoding, errors="ignore"
                    ) as f:
                        first_line = f.readline()
                    if "|" in first_line:
                        sep = "|"
                    elif "\t" in first_line:
                        sep = "\t"
                    elif ";" in first_line:
                        sep = ";"

                duckdb_sep = "\\t" if sep == "\t" else sep

                print(
                    _(
                        "Loading {} (encoding={}, sep='{}') into table {}...",
                        data_file.name,
                        duckdb_encoding,
                        duckdb_sep,
                        table_name,
                    )
                )

                # Create temp view to inspect columns
                conn.execute(
                    f"CREATE OR REPLACE TEMP VIEW temp_view AS SELECT * FROM read_csv('{escaped_path}', auto_detect=true, sep='{duckdb_sep}', encoding='{duckdb_encoding}')"
                )
            cols = conn.execute("DESCRIBE temp_view").fetchall()

            select_exprs = []
            seen_std_names = set()
            for c in cols:
                orig_col = c[0]
                std_name = reverse_map.get(orig_col.lower(), orig_col)

                # Failsafe: avoid duplicate column names in the same table
                if std_name.lower() in seen_std_names:
                    std_name = orig_col  # fallback to original if collision still happens somehow

                # If still duplicate (e.g. original name conflicts with a mapped name)
                counter = 1
                base_std_name = std_name
                while std_name.lower() in seen_std_names:
                    counter += 1
                    std_name = f"{base_std_name}_{counter}"

                seen_std_names.add(std_name.lower())

                escaped_orig = orig_col.replace('"', '""')
                escaped_std = std_name.replace('"', '""')
                if std_name != orig_col:
                    select_exprs.append(f'"{escaped_orig}" AS "{escaped_std}"')
                else:
                    select_exprs.append(f'"{escaped_orig}"')

            select_sql = ",\n                ".join(select_exprs)

            query = (
                f'CREATE TABLE "{table_name}" AS SELECT \n'
                f"                {select_sql} \n"
                f"            FROM temp_view"
            )
            conn.execute(query)
            conn.execute("DROP VIEW temp_view")
            if is_clinical:
                conn.unregister("_clinical_df")

            escaped_source = current_source.replace("'", "''")
            conn.execute(f"COMMENT ON TABLE \"{table_name}\" IS '{escaped_source}'")

            row_count_res = conn.execute(
                f'SELECT COUNT(*) FROM "{table_name}"'
            ).fetchone()
            row_count = row_count_res[0] if row_count_res else 0
            file_size = data_file.stat().st_size
            conn.execute(
                "DELETE FROM _t1d_ingestion_manifest WHERE source_file = ? OR table_name = ?",
                [current_source, table_name],
            )
            conn.execute(
                "INSERT INTO _t1d_ingestion_manifest VALUES (?, ?, ?, ?, CURRENT_TIMESTAMP, ?)",
                [current_source, table_name, current_hash, row_count, file_size],
            )

            loaded_tables.append(table_name)
        except Exception as e:
            print(_("Failed to load {}: {}", data_file.name, e))

    if loaded_tables:
        print(_("Successfully populated {} tables in {}.", len(loaded_tables), db_path))
    conn.close()


def get_database_schema(conn: duckdb.DuckDBPyConnection) -> str:
    """
    Extract the database schema as a text description, including sample data.

    Args:
    ----
        conn: The DuckDB connection.

    Returns:
    -------
        The database schema string.

    """
    tables = [
        t
        for t in conn.execute("SHOW TABLES").fetchall()
        if not t[0].startswith("_t1d_")
    ]
    schema_parts = []
    for (table_name,) in tables:
        columns = conn.execute(f'DESCRIBE "{table_name}"').fetchall()
        col_desc = ", ".join([f"{c[0]} ({c[1]})" for c in columns])

        try:
            sample = conn.execute(f'SELECT * FROM "{table_name}" LIMIT 1').fetchone()
            sample_desc = (
                f"Sample row: {sample}" if sample else "Sample row: (empty table)"
            )
        except Exception:
            sample_desc = "Sample row: (could not fetch)"

        schema_parts.append(f"Table: {table_name}\nColumns: {col_desc}\n{sample_desc}")
    return "\n\n".join(schema_parts)


def is_sql_query(query: str, conn: Optional[duckdb.DuckDBPyConnection] = None) -> bool:
    """
    Determine whether a user input string is a SQL query rather than natural language.

    Args:
    ----
        query: The raw query string entered by the user.
        conn: Optional DuckDB connection to assist in query syntax verification.

    Returns:
    -------
        True if the query appears to be a SQL statement, False if natural language.

    """
    # Strip SQL comments (-- ... and /* ... */) and whitespace
    cleaned = re.sub(r"--[^\n]*\n?", " ", query)
    cleaned = re.sub(r"/\*.*?\*/", " ", cleaned, flags=re.DOTALL).strip()
    if not cleaned:
        return False

    unwrapped = cleaned
    while unwrapped.startswith("(") and unwrapped.endswith(")"):
        unwrapped = unwrapped[1:-1].strip()

    if not unwrapped:
        return False

    first_word = unwrapped.split()[0].lower().rstrip(";").strip()
    sql_keywords = (
        "select",
        "with",
        "describe",
        "pragma",
        "explain",
        "summarize",
        "call",
        "values",
        "from",
        "table",
        "show",
        "create",
        "drop",
        "alter",
        "insert",
        "update",
        "delete",
        "copy",
        "attach",
        "detach",
        "begin",
        "commit",
        "rollback",
        "checkpoint",
    )
    if first_word not in sql_keywords:
        return False

    if first_word == "show":
        words = unwrapped.lower().split()
        nl_indicators = {
            "me",
            "the",
            "a",
            "an",
            "us",
            "our",
            "all",
            "how",
            "what",
            "where",
            "which",
            "any",
            "patients",
            "trials",
            "datasets",
        }
        if len(words) > 1 and words[1] in nl_indicators:
            return False

    temp_conn = conn or duckdb.connect(":memory:")
    try:
        if type(temp_conn).__name__ == "DuckDBPyConnection" and hasattr(
            temp_conn, "extract_statements"
        ):
            stmts = temp_conn.extract_statements(unwrapped)
            return len(stmts) > 0

        # Safe fallback without side-effects on connection
        if first_word in (
            "checkpoint",
            "attach",
            "detach",
            "call",
            "pragma",
            "create",
            "drop",
            "alter",
            "insert",
            "update",
            "delete",
        ):
            return True

        temp_conn.execute(f"EXPLAIN {unwrapped}")
        return True
    except duckdb.ParserException:
        return False
    except Exception:
        return True
    finally:
        if conn is None:
            temp_conn.close()


def extract_sql_from_response(full_response: str) -> str:
    """
    Extract SQL query from markdown code blocks or plain text in an LLM completion.

    Args:
    ----
        full_response: Raw text response from the LLM.

    Returns:
    -------
        str: Extracted SQL query string.

    """
    if not isinstance(full_response, str):
        full_response = str(full_response)

    blocks = re.findall(
        r"```(?:[a-zA-Z0-9_-]+[ \t]*\n)?(.*?)\n?```",
        full_response,
        re.DOTALL | re.IGNORECASE,
    )
    if blocks:
        sql_candidate = blocks[-1].strip()
        for block in reversed(blocks):
            clean = block.strip()
            if clean.lower().startswith(
                (
                    "select",
                    "with",
                    "show",
                    "describe",
                    "pragma",
                    "summarize",
                    "explain",
                )
            ):
                sql_candidate = clean
                break
        return str(sql_candidate)

    sql_query = full_response
    if sql_query.lower().startswith("sql\n"):
        sql_query = sql_query[4:]
    elif sql_query.lower().startswith("duckdb\n"):
        sql_query = sql_query[7:]
    return sql_query.replace("```", "").strip()


def handle_natural_language(
    conn: duckdb.DuckDBPyConnection,
    query: str,
    model: str = "gemma4",
    provider: Optional[str] = None,
) -> None:
    """
    Translate natural language to SQL using an LLM via any-llm and execute it.

    Args:
    ----
        conn: DB connection.
        query: The natural language query.
        model: Model name for translation (or 'provider/model').
        provider: Optional provider name ('ollama', 'openai', 'anthropic', 'google').

    """
    _ = get_translator()
    try:
        from any_llm import AnyLLM  # type: ignore[import-not-found]
    except ImportError:
        print(_("Error: any-llm-sdk[ollama] is not installed. Please install it."))
        return

    prov = provider.lower() if provider else "ollama"
    actual_model = model
    if "/" in model:
        prov, actual_model = model.split("/", 1)
        prov = prov.lower()

    from t1d_analytics.training_data import validate_provider_readiness

    try:
        validate_provider_readiness(prov)
    except RuntimeError as err:
        print(_("Failed to generate or execute query: {}", err))
        return

    print(_("Thinking..."))
    schema = get_database_schema(conn)

    prompt = f"""You are a DuckDB SQL expert. Given the following database schema for Type 1 Diabetes (T1D) clinical trial datasets:

{schema}

Context & Rules for T1D Analytics:
1. Terminology & Aliases:
   - "Lows" / "Hypoglycemia": Look for Time Below Range (TBR, <70 or <54 mg/dL), `pct_time_under_70`, `hypo`.
   - "Highs" / "Hyperglycemia": Look for Time Above Range (TAR, >180 or >250 mg/dL), `pct_time_over_180`, `hyper`.
   - "In-range" / "TIR": Look for Time in Range (70-180 mg/dL), `pct_time_in_range_70_180`, `tir`, `cgm_in_range`.
   - "A1C" / "HbA1c": Look for glycated hemoglobin, `a1c`, `hba1c`.
   - "Variability": Look for Coefficient of Variation (`cv`), Standard Deviation (`sd`).
   - "Average sugar": Look for Mean Glucose, `mean_cgm`, Glucose Management Indicator (`gmi`).
   - "Adverse Events": "DKA" (Diabetic Ketoacidosis), "SH" (Severe Hypoglycemia).
   - Demographics: "Duration" (years with diabetes), "Age", "BMI", "Gender/Sex", "Weight".
   - Treatment: "Pump" (CSII), "Injections/Shots" (MDI), "Loop" (AID/Closed-Loop), "Insulin" (TDD, Basal, Bolus).
   - Study Timepoints: "Baseline" (visit 0 or baseline flag), "End of trial/Follow-up" (max visit month).
2. SQL Generation Rules for Multi-Trial Data:
   - Trial Names: Acronyms (DCLP3, DCLP5, Pedap, etc.) usually correspond directly to table names.
   - Aggregating Combined Trials: To calculate metrics for multiple trials "together", "combined", or "overall", you MUST use a CTE with `UNION ALL` to combine raw rows from all requested tables *before* applying aggregate functions like `AVG()`.
   - Comparing Individual Trials: To calculate metrics for multiple trials "individually" or "each", use `UNION ALL` with an identifying literal string column (e.g., `SELECT 'DCLP3' as trial, AVG(...) ... UNION ALL SELECT 'DCLP5'...`).
   - Aligning UNION Columns: Clinical trial tables often have different schemas. When using `UNION ALL`, you MUST explicitly `SELECT` and alias only the specific columns needed (e.g. `age`, `tir`) to ensure they align perfectly in number, order, and type across all blocks. NEVER use `SELECT *` with `UNION ALL`.
   - Linking Data: If questions span patient demographics and CGM outcomes, use `JOIN` on the patient/subject ID columns (e.g., `pt_id`, `subject_id`).
3. Formatting & Logic:
   - Handle colloquial/informal language intuitively (e.g., "betwixt" -> `BETWEEN`, "kids/pediatric" -> `age < 18`, "adults" -> `age >= 18`).
   - Account for string vs. numeric data types; use `CAST(col AS FLOAT)` if computing averages on textual columns that contain numbers.
   - Ignore NULL values appropriately when calculating averages or sums.
4. Strict Schema Enforcement:
   - CRITICAL: You must ONLY use the exact table names and column names provided in the schema above.
   - Do NOT invent, guess, or hallucinate table names (e.g., do not use "Visit1" or "demographics" if they are not in the schema).
   - If a requested concept (like "demographics") is not a table, find the relevant columns within the existing tables (e.g., `visits` or `patients`).

Write a SQL query that answers the user's request. 
You may provide an explanation of your thought process as SQL comments (`--`) before the query.
Return ONLY the raw SQL query, with no markdown formatting and no code blocks.
The query should be a valid DuckDB SQL SELECT statement.

User request: {query}
"""
    try:
        if prov == "ollama":
            api_base = os.environ.get("OLLAMA_HOST", "http://127.0.0.1:11434")
            if not api_base.startswith("http://") and not api_base.startswith(
                "https://"
            ):
                api_base = f"http://{api_base}"

        llm = AnyLLM.create(prov)
        response = llm.completion(
            model=actual_model,
            messages=[{"role": "user", "content": prompt}],
        )
        full_response = response.choices[0].message.content.strip()
        sql_query = extract_sql_from_response(full_response)

        print(_("Generated SQL: \n{}\n", sql_query))
        print(_("Executing...\n"))

        result = conn.sql(sql_query)
        if result:
            result.show()

    except Exception as e:
        print(_("Failed to generate or execute query: {}", e))


def run_query_repl(
    db_path: str, model: str = "gemma4", provider: Optional[str] = None
) -> None:
    """
    Run the interactive query interface.

    Args:
    ----
        db_path: Path to DB.
        model: LLM model name for query translation (or 'provider/model').
        provider: Optional LLM provider name ('ollama', 'openai', 'anthropic', 'google').

    """
    _ = get_translator()
    if not Path(db_path).exists():
        print(_("Database {} does not exist.", db_path))
        print(_("Please run the 'load' command first to populate the database."))
        return

    db_health = check_database_health(db_path)
    if db_health.status_code == DatabaseStatusCode.EMPTY_DB:
        print(_("Warning: Database {} contains 0 tables.", db_path))
        print(_("Please run the 'load' command first to populate the database."))
    elif db_health.status_code == DatabaseStatusCode.MISSING_INITIAL_DATA:
        print(
            _(
                "Warning: Database {} lacks standard initial clinical trial datasets.",
                db_path,
            )
        )

    print(_("Connecting to DuckDB at {}...", db_path))
    conn = duckdb.connect(db_path, read_only=True)

    print("\n" + "=" * 50)
    print(_("Welcome to T1D Analytics Interface!"))
    print(_("You can enter:"))
    print(
        _("  - Standard SQL queries (starting with SELECT, WITH, SHOW, DESCRIBE, etc.)")
    )
    print(_("  - Natural language queries (will be translated to SQL via LLM)"))
    print(_("  - 'exit' or 'quit' to close."))
    print("=" * 50 + "\n")

    while True:
        try:
            user_input = input("query> ").strip()
        except (EOFError, KeyboardInterrupt):
            print("\n" + _("Exiting."))
            break

        if not user_input:
            continue

        lower_input = user_input.lower()
        if lower_input in ("exit", "quit"):
            print(_("Exiting."))
            break

        if is_sql_query(user_input, conn=conn):
            try:
                result = conn.sql(user_input)
                # Show results nicely
                if result:
                    result.show()
            except Exception as e:
                print(_("SQL Error: {}", e))
        else:
            handle_natural_language(conn, user_input, model=model, provider=provider)

    conn.close()


def profile_table(db_path: str, table_name: str) -> Dict[str, Any]:
    """
    Profile a table to calculate data quality, completeness, and column statistics.

    Args:
    ----
        db_path: Path to DuckDB file.
        table_name: Table name to profile.

    Returns:
    -------
        Dict[str, Any]: Profiling summary containing total_rows, columns stats, null percentages.

    Raises:
    ------
        ValueError: If table name is invalid or table does not exist.
        FileNotFoundError: If database file does not exist.

    """
    if not table_name.isidentifier():
        raise ValueError(f"Invalid table identifier: {table_name}")

    if not Path(db_path).exists():
        raise FileNotFoundError(f"Database file not found: {db_path}")

    conn = duckdb.connect(db_path, read_only=True)
    try:
        tables = [r[0] for r in conn.execute("SHOW TABLES").fetchall()]
        if table_name not in tables:
            raise ValueError(f"Table '{table_name}' does not exist in database.")

        cols = conn.execute(f'DESCRIBE "{table_name}"').fetchall()
        column_profiles: List[Dict[str, Any]] = []

        if not cols:
            count_row = conn.execute(f'SELECT COUNT(*) FROM "{table_name}"').fetchone()
            empty_total_rows = count_row[0] if count_row else 0
            return {
                "table_name": table_name,
                "total_rows": empty_total_rows,
                "column_count": 0,
                "columns": [],
            }

        select_exprs: List[str] = ['COUNT(*) AS "_total_rows"']
        col_meta: List[Dict[str, Any]] = []

        for idx, col_info in enumerate(cols):
            col_name = str(col_info[0])
            col_type = str(col_info[1])
            escaped_col = col_name.replace('"', '""')

            select_exprs.append(
                f'SUM(CASE WHEN "{escaped_col}" IS NULL THEN 1 ELSE 0 END) AS "null_{idx}"'
            )
            select_exprs.append(f'COUNT(DISTINCT "{escaped_col}") AS "dist_{idx}"')

            type_upper = col_type.upper()
            is_numeric = any(
                t in type_upper
                for t in (
                    "INT",
                    "FLOAT",
                    "DOUBLE",
                    "DECIMAL",
                    "NUMERIC",
                    "HUGEINT",
                )
            )
            if is_numeric:
                select_exprs.append(f'MIN("{escaped_col}") AS "min_{idx}"')
                select_exprs.append(f'MAX("{escaped_col}") AS "max_{idx}"')
                select_exprs.append(f'AVG("{escaped_col}") AS "avg_{idx}"')

            col_meta.append(
                {
                    "column_name": col_name,
                    "data_type": col_type,
                    "is_numeric": is_numeric,
                }
            )

        escaped_tbl = table_name.replace('"', '""')
        agg_sql = f'SELECT {", ".join(select_exprs)} FROM "{escaped_tbl}"'
        agg_res = conn.execute(agg_sql).fetchone()

        total_rows: int = int(agg_res[0]) if (agg_res and agg_res[0] is not None) else 0

        cursor = 1
        for meta in col_meta:
            null_val = agg_res[cursor] if agg_res else 0
            null_count: int = int(null_val) if null_val is not None else 0
            cursor += 1

            dist_val = agg_res[cursor] if agg_res else 0
            distinct_count: int = int(dist_val) if dist_val is not None else 0
            cursor += 1

            null_pct = (null_count / total_rows * 100.0) if total_rows > 0 else 0.0

            col_stat: Dict[str, Any] = {
                "column_name": meta["column_name"],
                "data_type": meta["data_type"],
                "null_count": null_count,
                "null_percentage": round(null_pct, 2),
                "distinct_count": distinct_count,
            }

            if meta["is_numeric"]:
                min_val = agg_res[cursor] if agg_res else None
                cursor += 1
                max_val = agg_res[cursor] if agg_res else None
                cursor += 1
                avg_val = agg_res[cursor] if agg_res else None
                cursor += 1

                if min_val is not None and max_val is not None:
                    col_stat["min"] = min_val
                    col_stat["max"] = max_val
                    col_stat["mean"] = (
                        round(float(avg_val), 2) if avg_val is not None else None
                    )
                else:
                    col_stat["min"] = None
                    col_stat["max"] = None
                    col_stat["mean"] = None

            column_profiles.append(col_stat)

        return {
            "table_name": table_name,
            "total_rows": total_rows,
            "column_count": len(column_profiles),
            "columns": column_profiles,
        }
    finally:
        conn.close()


def export_table(
    db_path: str,
    table_name: str,
    output_path: Union[str, Path],
    export_format: str = "excel",
) -> None:
    """
    Export a DuckDB table to Excel (.xlsx), Parquet, or CSV format with security sanitization.

    Args:
    ----
        db_path: Path to the DuckDB database file.
        table_name: Table identifier to export.
        output_path: Target filesystem path for the exported file.
        export_format: Format string, one of 'excel', 'parquet', or 'csv'.

    Raises:
    ------
        FileNotFoundError: If the database file does not exist.
        ValueError: If the table identifier is invalid, the table does not exist, or format is unsupported.

    """
    if not table_name.isidentifier():
        raise ValueError(f"Invalid table identifier: {table_name}")

    if not Path(db_path).exists():
        raise FileNotFoundError(f"Database file not found: {db_path}")

    fmt = export_format.lower().strip()
    if fmt not in ("excel", "parquet", "csv"):
        raise ValueError(
            f"Unsupported export format: {export_format}. Choose 'excel', 'parquet', or 'csv'."
        )

    target_path = Path(output_path)
    target_path.parent.mkdir(parents=True, exist_ok=True)

    conn = duckdb.connect(db_path, read_only=True)
    try:
        tables = [r[0] for r in conn.execute("SHOW TABLES").fetchall()]
        if table_name not in tables:
            raise ValueError(f"Table '{table_name}' does not exist in database.")

        escaped_tbl = table_name.replace('"', '""')

        if fmt == "parquet":
            escaped_out = str(target_path.resolve()).replace("'", "''")
            conn.execute(f"COPY \"{escaped_tbl}\" TO '{escaped_out}' (FORMAT PARQUET)")
        elif fmt == "csv":
            escaped_out = str(target_path.resolve()).replace("'", "''")
            conn.execute(
                f"COPY \"{escaped_tbl}\" TO '{escaped_out}' (FORMAT CSV, HEADER TRUE)"
            )
        else:
            from openpyxl import Workbook

            result = conn.execute(f'SELECT * FROM "{escaped_tbl}"')
            columns = (
                [desc[0] for desc in result.description] if result.description else []
            )

            wb = Workbook(write_only=True)
            ws = wb.create_sheet(title=table_name[:31])
            ws.append(columns)

            while True:
                batch = result.fetchmany(5000)
                if not batch:
                    break
                for row in batch:
                    formatted_row: List[Any] = []
                    for val in row:
                        if val is None:
                            formatted_row.append(None)
                        elif isinstance(val, (int, float, bool)):
                            formatted_row.append(val)
                        else:
                            s_val = str(val)
                            if s_val and s_val[0] in ("=", "+", "-", "@", "\t", "\r"):
                                is_numeric = False
                                if s_val[0] in ("+", "-"):
                                    try:
                                        float(s_val)
                                        is_numeric = True
                                    except ValueError:
                                        is_numeric = False
                                if not is_numeric:
                                    s_val = f"'{s_val}"
                            formatted_row.append(s_val)
                    ws.append(formatted_row)

            wb.save(str(target_path))
    finally:
        conn.close()
