"""Tests for FastAPI schema, table preview, profiling, and CSV/JSON/Excel/Parquet export endpoints."""

import io
import json
import typing
from pathlib import Path
from unittest.mock import MagicMock, patch

import duckdb
import pytest
from fastapi.testclient import TestClient

from t1d_analytics.api import app, execute_sql
from t1d_analytics.sql_utils import SqlValue

client = TestClient(app)


def test_get_schema_success(mock_db: str) -> None:
    """Test schema endpoint returns structured schema from a real DuckDB database."""
    response = client.get(f"/api/schema?db_path={mock_db}")
    assert response.status_code == 200
    data = response.json()
    assert len(data["tables"]) == 1
    assert data["tables"][0]["name"] == "users"
    assert data["tables"][0]["columns"][0]["name"] == "id"
    assert data["tables"][0]["columns"][0]["type"] == "INTEGER"
    assert data["tables"][0]["columns"][1]["name"] == "name"
    assert data["tables"][0]["columns"][1]["type"] == "VARCHAR"


@patch("t1d_analytics.api.duckdb.connect")
def test_get_schema_error(mock_connect: MagicMock) -> None:
    """Test schema endpoint handles errors gracefully."""
    mock_connect.side_effect = Exception("DB connection failed")

    response = client.get("/api/schema")
    assert response.status_code == 200
    data = response.json()
    assert len(data["tables"]) == 0


def test_get_schema_with_db_path(mock_db: str) -> None:
    """Test get_schema respects db_path query parameter."""
    response = client.get(f"/api/schema?db_path={mock_db}")
    assert response.status_code == 200
    data = response.json()
    assert len(data["tables"]) >= 1
    table_names = [t["name"] for t in data["tables"]]
    assert "users" in table_names


def test_get_table_data(mock_db: str, monkeypatch: pytest.MonkeyPatch) -> None:
    """Test fetching valid table data with pagination, sorting, and searching."""
    monkeypatch.setenv("T1D_DB_PATH", mock_db)
    response = client.get("/api/table/users")
    assert response.status_code == 200
    data = response.json()
    assert "rows" in data
    assert len(data["rows"]) == 2
    assert data["rows"][0]["id"] == 1
    assert data["rows"][0]["name"] == "Alice"

    resp_desc = client.get("/api/table/users?sort_by=name&order=desc")
    assert resp_desc.status_code == 200
    assert resp_desc.json()["rows"][0]["name"] == "Bob"

    resp_bad_order = client.get("/api/table/users?order=invalid")
    assert resp_bad_order.status_code == 400
    assert resp_bad_order.json()["detail"]["error_code"] == "backend.invalidSortOrder"

    resp_bad_col = client.get("/api/table/users?sort_by=nonexistent")
    assert resp_bad_col.status_code == 400
    assert resp_bad_col.json()["detail"]["error_code"] == "backend.invalidSortColumn"

    resp_search = client.get("/api/table/users?search=alice")
    assert resp_search.status_code == 200
    assert len(resp_search.json()["rows"]) == 1
    assert resp_search.json()["rows"][0]["name"] == "Alice"
    assert resp_search.json()["total_count"] == 1

    resp_no_match = client.get("/api/table/users?search=nonexistent")
    assert resp_no_match.status_code == 200
    assert len(resp_no_match.json()["rows"]) == 0
    assert resp_no_match.json()["total_count"] == 0


def test_get_table_data_invalid_table(
    mock_db: str, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Test fetching data from a table that doesn't exist."""
    monkeypatch.setenv("T1D_DB_PATH", mock_db)
    response = client.get("/api/table/non_existent_table")
    assert response.status_code == 404


def test_get_table_data_sql_injection(
    mock_db: str, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Test protection against basic SQL injection in table name."""
    monkeypatch.setenv("T1D_DB_PATH", mock_db)
    response = client.get("/api/table/invalid table name;")
    assert response.status_code == 400


def test_get_table_data_exception(
    mock_db: str, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Test server error handling when DB operations fail."""
    monkeypatch.setenv("T1D_DB_PATH", mock_db)
    with patch("duckdb.connect", side_effect=Exception("DB Failure")):
        response = client.get("/api/table/users")
        assert response.status_code == 500
        assert response.json()["detail"]["error_code"] == "backend.serverError"
        assert "DB Failure" in response.json()["detail"]["params"]["error"]


def test_get_table_data_pagination_metadata(mock_db: str) -> None:
    """Test get_table_data returns total_count, page, and total_pages."""
    response = client.get(f"/api/table/users?limit=1&offset=0&db_path={mock_db}")
    assert response.status_code == 200
    data = response.json()
    assert data["total_count"] == 2
    assert data["page"] == 1
    assert data["total_pages"] == 2
    assert len(data["rows"]) == 1


def test_get_table_data_negative_pagination(mock_db: str) -> None:
    """Test get_table_data rejects negative or zero limit and negative offset."""
    resp_neg_lim = client.get("/api/table/users?limit=-1&db_path=" + mock_db)
    assert resp_neg_lim.status_code == 400
    assert resp_neg_lim.json()["detail"]["error_code"] == "backend.invalidPagination"

    resp_zero_lim = client.get("/api/table/users?limit=0&db_path=" + mock_db)
    assert resp_zero_lim.status_code == 400
    assert resp_zero_lim.json()["detail"]["error_code"] == "backend.invalidPagination"

    resp_neg_off = client.get("/api/table/users?offset=-5&db_path=" + mock_db)
    assert resp_neg_off.status_code == 400
    assert resp_neg_off.json()["detail"]["error_code"] == "backend.invalidPagination"


def test_get_table_data_missing_db_and_db_error(
    mock_db: str, mocker: MagicMock
) -> None:
    """Test get_table_data 404 on missing DB file and connection error."""
    resp_no_db = client.get("/api/table/users?db_path=nonexistent.duckdb")
    assert resp_no_db.status_code == 404

    mocker.patch("duckdb.connect", side_effect=Exception("Cannot open file"))
    resp_open_err = client.get(f"/api/table/users?db_path={mock_db}")
    assert resp_open_err.status_code == 404


def test_get_table_data_zero_columns(tmp_path: Path) -> None:
    """Test get_table_data gracefully handles tables when table_info reports 0 columns."""
    db_file = tmp_path / "zero_cols.duckdb"
    conn = duckdb.connect(str(db_file))
    conn.execute("CREATE TABLE dummy_tbl (id INT)")
    conn.close()

    orig_connect = duckdb.connect

    class MockConn:
        """Mock DuckDB connection that returns empty table_info."""

        def __init__(self, real_conn: duckdb.DuckDBPyConnection) -> None:
            """Store real connection."""
            self._real = real_conn

        def execute(
            self, query: str, params: typing.Optional[list[SqlValue]] = None
        ) -> typing.Union[MagicMock, duckdb.DuckDBPyConnection]:
            """Return empty metadata on PRAGMA table_info/DESCRIBE."""
            if "PRAGMA table_info" in query or "DESCRIBE" in query:
                mock_res = MagicMock()
                mock_res.fetchall.return_value = []
                mock_res.description = []
                return mock_res
            if params is not None:
                return self._real.execute(query, params)
            return self._real.execute(query)

        def close(self) -> None:
            """Close underlying real connection."""
            self._real.close()

    def mock_conn_factory(path: str, read_only: bool = False) -> MockConn:
        """Create a MockConn wrapping a real DuckDB connection."""
        real = orig_connect(path, read_only=read_only)
        return MockConn(real)

    with patch("duckdb.connect", side_effect=mock_conn_factory):
        resp = client.get(
            "/api/table/dummy_tbl",
            params={"db_path": str(db_file)},
        )
        assert resp.status_code == 200
        data = resp.json()
        assert data["rows"] == []
        assert data["total_count"] == 0


def test_sanitize_sql_value() -> None:
    """Test sanitize_sql_value with diverse Python and DuckDB data types."""
    import datetime
    import decimal
    import uuid

    from t1d_analytics.api import sanitize_sql_value

    assert sanitize_sql_value(None) is None
    assert sanitize_sql_value(datetime.date(2023, 1, 15)) == "2023-01-15"
    assert (
        sanitize_sql_value(datetime.datetime(2023, 1, 15, 12, 30, 45))
        == "2023-01-15T12:30:45"
    )
    assert sanitize_sql_value(datetime.time(8, 30)) == "08:30:00"
    assert sanitize_sql_value(decimal.Decimal("123.45")) == 123.45
    u = uuid.uuid4()
    assert sanitize_sql_value(u) == str(u)
    assert sanitize_sql_value(float("nan")) is None
    assert sanitize_sql_value(float("inf")) is None
    assert sanitize_sql_value(float("-inf")) is None
    assert sanitize_sql_value(12.34) == 12.34
    assert sanitize_sql_value(42) == 42
    assert sanitize_sql_value(True) is True
    assert sanitize_sql_value("test string") == "test string"
    assert sanitize_sql_value([1, 2, 3]) == "[1, 2, 3]"


def test_execute_sql_complex_types(tmp_path: Path) -> None:
    """Test execute_sql and TableDataResponse with dates, timestamps, decimals, and nan."""
    db_path = str(tmp_path / "complex.duckdb")
    conn = duckdb.connect(db_path)
    conn.execute(
        """
        CREATE TABLE clinical (
            id INTEGER,
            reading_date DATE,
            reading_ts TIMESTAMP,
            val DECIMAL(5, 2),
            nan_val DOUBLE
        )
        """
    )
    conn.execute(
        """
        INSERT INTO clinical VALUES (
            1,
            '2023-05-10'::DATE,
            '2023-05-10 14:20:00'::TIMESTAMP,
            120.50,
            'nan'::DOUBLE
        )
        """
    )
    conn.close()

    res = execute_sql(db_path, "SELECT * FROM clinical")
    assert len(res) == 1
    row = res[0]
    assert row["id"] == 1
    assert row["reading_date"] == "2023-05-10"
    assert row["reading_ts"] == "2023-05-10T14:20:00"
    assert row["val"] == 120.5
    assert row["nan_val"] is None

    table_resp = client.get(f"/api/table/clinical?db_path={db_path}")
    assert table_resp.status_code == 200
    t_data = table_resp.json()
    assert t_data["total_count"] == 1
    assert t_data["rows"][0]["reading_date"] == "2023-05-10"

    chat_resp = client.post(
        "/api/chat",
        json={
            "message": "SELECT * FROM clinical",
            "model": "sql",
            "db_path": db_path,
        },
    )
    assert chat_resp.status_code == 200
    c_data = chat_resp.json()
    assert c_data["sqlResult"][0]["val"] == 120.5


def test_get_table_profile_endpoint(mock_db: str, mocker: MagicMock) -> None:
    """Test /api/table/{table_name}/profile endpoint success and error conditions."""
    resp = client.get(f"/api/table/users/profile?db_path={mock_db}")
    assert resp.status_code == 200
    data = resp.json()
    assert data["table_name"] == "users"
    assert data["total_rows"] == 2

    resp_missing_db = client.get("/api/table/users/profile?db_path=missing_file.duckdb")
    assert resp_missing_db.status_code == 404

    resp_bad_id = client.get(f"/api/table/users;drop/profile?db_path={mock_db}")
    assert resp_bad_id.status_code == 400

    resp_no_tbl = client.get(f"/api/table/missing_tbl/profile?db_path={mock_db}")
    assert resp_no_tbl.status_code == 404

    mocker.patch(
        "t1d_analytics.api.profile_table",
        side_effect=FileNotFoundError("not found"),
    )
    resp_fnf = client.get(f"/api/table/users/profile?db_path={mock_db}")
    assert resp_fnf.status_code == 404

    mocker.patch(
        "t1d_analytics.api.profile_table",
        side_effect=Exception("Database crash"),
    )
    resp_err = client.get(f"/api/table/users/profile?db_path={mock_db}")
    assert resp_err.status_code == 500


def test_export_excel_endpoint(mock_db: str, mocker: MagicMock) -> None:
    """Test /api/export/excel endpoint streaming and error branches."""
    import openpyxl

    resp = client.get(f"/api/export/excel?table_name=users&db_path={mock_db}")
    assert resp.status_code == 200
    assert (
        resp.headers["content-type"]
        == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "users_export.xlsx" in resp.headers["content-disposition"]
    wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    ws = wb.active
    assert ws is not None
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0] == ("id", "name")
    assert any("Alice" in str(r) for r in rows)

    conn = duckdb.connect(mock_db)
    conn.execute(
        "CREATE TABLE formulas (id INT, val_neg DOUBLE, val_pos VARCHAR, formula1 VARCHAR, formula2 VARCHAR, null_col VARCHAR)"
    )
    conn.execute(
        "INSERT INTO formulas VALUES (1, -5.2, '+10', '=cmd|'' /C calc''!A0', '-cmd', NULL)"
    )
    conn.close()

    resp_f = client.get(f"/api/export/excel?table_name=formulas&db_path={mock_db}")
    assert resp_f.status_code == 200
    wb_f = openpyxl.load_workbook(io.BytesIO(resp_f.content))
    ws_f = wb_f.active
    assert ws_f is not None
    rows_f = list(ws_f.iter_rows(values_only=True))
    data_row = rows_f[1]
    assert data_row[1] == -5.2
    assert data_row[2] == "+10"
    assert data_row[3] == "'=cmd|' /C calc'!A0"
    assert data_row[4] == "'-cmd"
    assert data_row[5] is None

    resp_no_db = client.get("/api/export/excel?table_name=users&db_path=no_db.duckdb")
    assert resp_no_db.status_code == 404

    resp_inv = client.get(f"/api/export/excel?table_name=users;drop&db_path={mock_db}")
    assert resp_inv.status_code == 400

    resp_no_tbl = client.get(f"/api/export/excel?table_name=missing&db_path={mock_db}")
    assert resp_no_tbl.status_code == 404

    mocker.patch("duckdb.connect", side_effect=Exception("IO Failure"))
    resp_500 = client.get(f"/api/export/excel?table_name=users&db_path={mock_db}")
    assert resp_500.status_code == 500

    mock_conn_exc = mocker.MagicMock()
    mock_conn_exc.execute.side_effect = [
        mocker.MagicMock(fetchall=lambda: [("users",)]),
        Exception("Execution failure"),
    ]
    mocker.patch("duckdb.connect", return_value=mock_conn_exc)
    resp_500b = client.get(f"/api/export/excel?table_name=users&db_path={mock_db}")
    assert resp_500b.status_code == 500


def test_export_parquet_endpoint(mock_db: str, mocker: MagicMock) -> None:
    """Test /api/export/parquet endpoint for data export in columnar parquet format."""
    resp = client.get(f"/api/export/parquet?table_name=users&db_path={mock_db}")
    assert resp.status_code == 200
    assert resp.headers["content-type"] == "application/vnd.apache.parquet"
    assert "users_export.parquet" in resp.headers["content-disposition"]
    assert len(resp.content) > 0

    resp_no_db = client.get("/api/export/parquet?table_name=users&db_path=no_db.duckdb")
    assert resp_no_db.status_code == 404

    resp_inv = client.get(
        f"/api/export/parquet?table_name=users;drop&db_path={mock_db}"
    )
    assert resp_inv.status_code == 400

    resp_no_tbl = client.get(
        f"/api/export/parquet?table_name=missing&db_path={mock_db}"
    )
    assert resp_no_tbl.status_code == 404

    mocker.patch("duckdb.connect", side_effect=Exception("Parquet export failure"))
    resp_500 = client.get(f"/api/export/parquet?table_name=users&db_path={mock_db}")
    assert resp_500.status_code == 500

    mock_conn_p = mocker.MagicMock()
    mock_conn_p.execute.side_effect = [
        mocker.MagicMock(fetchall=lambda: [("users",)]),
        Exception("Disk error during COPY"),
    ]
    mocker.patch("duckdb.connect", return_value=mock_conn_p)
    resp_500b = client.get(f"/api/export/parquet?table_name=users&db_path={mock_db}")
    assert resp_500b.status_code == 500


def test_export_csv_endpoint(mock_db: str, mocker: MagicMock) -> None:
    """Test /api/export/csv endpoint streaming, formula escaping, and error handling."""
    resp = client.get(f"/api/export/csv?table_name=users&db_path={mock_db}")
    assert resp.status_code == 200
    assert "text/csv" in resp.headers["content-type"]
    assert "users_export.csv" in resp.headers["content-disposition"]
    text = resp.content.decode("utf-8")
    assert text.startswith(chr(0xFEFF))
    assert '"id","name"' in text or '"id"' in text

    conn = duckdb.connect(mock_db)
    conn.execute("CREATE TABLE csv_formulas (id INT, val VARCHAR, num VARCHAR)")
    conn.execute(
        "INSERT INTO csv_formulas VALUES (1, '=CMD', NULL), (2, '@SUM', '-99.9'), (3, '+CMD', '-CALC')"
    )
    conn.execute("CREATE TABLE multi_chunk (id INT, txt VARCHAR)")
    conn.execute(
        "INSERT INTO multi_chunk SELECT i, repeat('X', 120) FROM range(800) tbl(i)"
    )
    conn.close()
    resp_f = client.get(f"/api/export/csv?table_name=csv_formulas&db_path={mock_db}")
    assert resp_f.status_code == 200
    text_f = resp_f.content.decode("utf-8")
    assert "'=CMD" in text_f
    assert '""' in text_f
    assert "'+42.5" not in text_f
    assert "'+99.9" not in text_f

    resp_multi = client.get(f"/api/export/csv?table_name=multi_chunk&db_path={mock_db}")
    assert resp_multi.status_code == 200
    assert len(resp_multi.content) > 70000

    resp_no_db = client.get("/api/export/csv?table_name=users&db_path=no_db.duckdb")
    assert resp_no_db.status_code == 404

    resp_inv = client.get(f"/api/export/csv?table_name=users;drop&db_path={mock_db}")
    assert resp_inv.status_code == 400

    resp_no_tbl = client.get(f"/api/export/csv?table_name=missing&db_path={mock_db}")
    assert resp_no_tbl.status_code == 404

    mocker.patch("duckdb.connect", side_effect=Exception("CSV export failure"))
    resp_500 = client.get(f"/api/export/csv?table_name=users&db_path={mock_db}")
    assert resp_500.status_code == 500

    mock_conn_open = mocker.MagicMock()
    mock_conn_open.execute.side_effect = Exception("Show tables error")
    mocker.patch("duckdb.connect", return_value=mock_conn_open)
    resp_500b = client.get(f"/api/export/csv?table_name=users&db_path={mock_db}")
    assert resp_500b.status_code == 500
    mock_conn_open.close.assert_called()

    mock_conn = mocker.MagicMock()
    mock_res = mocker.MagicMock()
    mock_res.description = [("id",)]
    mock_res.fetchmany.side_effect = Exception("Read failure")
    mock_conn.execute.side_effect = [
        mocker.MagicMock(fetchall=lambda: [("users",)]),
        mock_res,
    ]
    mocker.patch("duckdb.connect", return_value=mock_conn)
    resp_500c = client.get(f"/api/export/csv?table_name=users&db_path={mock_db}")
    assert resp_500c.status_code == 500


def test_export_json_endpoint(mock_db: str, mocker: MagicMock) -> None:
    """Test /api/export/json endpoint streaming JSON array and error handling."""
    resp = client.get(f"/api/export/json?table_name=users&db_path={mock_db}")
    assert resp.status_code == 200
    assert "application/json" in resp.headers["content-type"]
    assert "users_export.json" in resp.headers["content-disposition"]
    data = json.loads(resp.content.decode("utf-8"))
    assert isinstance(data, list)
    assert len(data) > 0

    conn = duckdb.connect(mock_db)
    conn.execute("CREATE TABLE multi_chunk (id INT, txt VARCHAR)")
    conn.execute(
        "INSERT INTO multi_chunk SELECT i, repeat('X', 120) FROM range(800) tbl(i)"
    )
    conn.close()
    resp_multi_j = client.get(
        f"/api/export/json?table_name=multi_chunk&db_path={mock_db}"
    )
    assert resp_multi_j.status_code == 200
    assert len(resp_multi_j.content) > 70000

    resp_no_db = client.get("/api/export/json?table_name=users&db_path=no_db.duckdb")
    assert resp_no_db.status_code == 404

    resp_inv = client.get(f"/api/export/json?table_name=users;drop&db_path={mock_db}")
    assert resp_inv.status_code == 400

    resp_no_tbl = client.get(f"/api/export/json?table_name=missing&db_path={mock_db}")
    assert resp_no_tbl.status_code == 404

    mocker.patch("duckdb.connect", side_effect=Exception("JSON export failure"))
    resp_500 = client.get(f"/api/export/json?table_name=users&db_path={mock_db}")
    assert resp_500.status_code == 500

    mock_conn_open = mocker.MagicMock()
    mock_conn_open.execute.side_effect = Exception("JSON show tables error")
    mocker.patch("duckdb.connect", return_value=mock_conn_open)
    resp_500b = client.get(f"/api/export/json?table_name=users&db_path={mock_db}")
    assert resp_500b.status_code == 500
    mock_conn_open.close.assert_called()

    mock_conn = mocker.MagicMock()
    mock_res = mocker.MagicMock()
    mock_res.description = [("id",)]
    mock_res.fetchmany.side_effect = Exception("JSON read failure")
    mock_conn.execute.side_effect = [
        mocker.MagicMock(fetchall=lambda: [("users",)]),
        mock_res,
    ]
    mocker.patch("duckdb.connect", return_value=mock_conn)
    resp_500c = client.get(f"/api/export/json?table_name=users&db_path={mock_db}")
    assert resp_500c.status_code == 500


def test_export_excel_with_null_and_formula(tmp_path: Path) -> None:
    """Test /api/export/excel with null values and formula injection prefix."""
    import openpyxl

    db_path = str(tmp_path / "formula.duckdb")
    conn = duckdb.connect(db_path)
    conn.execute("CREATE TABLE items (id INTEGER, formula VARCHAR, note VARCHAR)")
    conn.execute("INSERT INTO items VALUES (1, '=1+1', NULL)")
    conn.close()

    resp = client.get(f"/api/export/excel?table_name=items&db_path={db_path}")
    assert resp.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    ws = wb.active
    assert ws is not None
    rows = list(ws.iter_rows(values_only=True))
    assert rows[1][1] == "'=1+1"
    assert rows[1][2] is None


def test_api_special_characters_and_quotes(tmp_path: Path) -> None:
    """Test get_schema, get_table_data, and export_excel with quotes and special characters."""
    import openpyxl

    db = tmp_path / "api_special.duckdb"
    conn = duckdb.connect(str(db))
    conn.execute(
        'CREATE TABLE "users_special" ("id" INT, "Glucose (mg/dL)" FLOAT, "col""quoted""" VARCHAR)'
    )
    conn.execute(
        """INSERT INTO "users_special" VALUES (1, 105.0, 'quote "inside" text'), (2, 120.0, 'normal')"""
    )
    conn.close()

    resp_schema = client.get(f"/api/schema?db_path={db}")
    assert resp_schema.status_code == 200
    assert any(t["name"] == "users_special" for t in resp_schema.json()["tables"])

    resp_data = client.get(f"/api/table/users_special?db_path={db}")
    assert resp_data.status_code == 200
    assert len(resp_data.json()["rows"]) == 2

    resp_export = client.get(f"/api/export/excel?table_name=users_special&db_path={db}")
    assert resp_export.status_code == 200
    wb_sp = openpyxl.load_workbook(io.BytesIO(resp_export.content))
    ws_sp = wb_sp.active
    assert ws_sp is not None
    rows_sp = list(ws_sp.iter_rows(values_only=True))
    assert rows_sp[0] == ("id", "Glucose (mg/dL)", 'col"quoted"')
    assert rows_sp[1] == (1, 105.0, 'quote "inside" text')
